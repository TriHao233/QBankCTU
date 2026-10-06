#!/usr/bin/env python3
"""Sinh tài liệu DOCX: các luồng nghiệp vụ và phân quyền QBankCTU."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_run_font(run, *, bold=False, size=11, color=None, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if color:
        run.font.color.rgb = color


def add_heading_custom(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, bold=True, size={1: 16, 2: 14, 3: 12}.get(level, 11))
    return p


def add_para(doc, text, *, bold=False, italic=False, size=11, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, bold=bold, italic=italic, size=size)
    return p


def add_bullet(doc, text, *, level=0):
    p = doc.add_paragraph(style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Cm(0.75 * (level + 1))
    p.clear()
    run = p.add_run(text)
    set_run_font(run, size=11)
    p.paragraph_format.space_after = Pt(3)
    return p


def set_cell_text(cell, text, *, bold=False, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, bold=bold, size=10)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, center=True)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            set_cell_text(table.rows[r_idx + 1].cells[c_idx], val)
    doc.add_paragraph()
    return table


def build_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)

    # ========== TITLE ==========
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(
        "TÀI LIỆU LUỒNG NGHIỆP VỤ VÀ PHÂN QUYỀN HỆ THỐNG"
    )
    set_run_font(run, bold=True, size=18, color=RGBColor(0x1A, 0x3A, 0x5C))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("QBankCTU — Ngân hàng câu hỏi tích hợp LLM & RAG")
    set_run_font(run, bold=True, size=13, color=RGBColor(0x2C, 0x5F, 0x8A))

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(
        "Tài liệu được tổng hợp từ codebase hiện tại "
        "(backend/core/dependencies.py, modules/*, frontend/src/auth/permissions.js).\n"
        "Nguyên tắc cốt lõi: Human-in-the-loop — AI chỉ tạo bản nháp và đánh giá; "
        "con người quyết định duyệt cuối cùng."
    )
    set_run_font(run, italic=True, size=10)

    # ========== MỤC LỤC ==========
    add_heading_custom(doc, "Mục lục", 1)
    toc_items = [
        "1. Tổng quan hệ thống và mô hình phân quyền",
        "2. Ba vai trò người dùng và quyền mặc định",
        "3. Danh mục quyền (permission catalog)",
        "4. Luồng xác thực, đăng ký và kích hoạt tài khoản",
        "5. Luồng quản lý tài liệu (Upload → OCR → Chunk → RAG)",
        "6. Luồng sinh câu hỏi bằng AI",
        "7. Luồng soạn thảo, gửi duyệt và quản lý ngân hàng câu hỏi",
        "8. Luồng đánh giá AI (AI Evaluation)",
        "9. Luồng kiểm duyệt người (Human Review)",
        "10. Luồng xuất bản Moodle",
        "11. Luồng làm đề thi",
        "12. Luồng Catalog (học phần, model, prompt, policy)",
        "13. Luồng Dictionary (từ điển từ khóa)",
        "14. Luồng quản lý người dùng (Admin)",
        "15. Luồng thông báo (Notifications)",
        "16. Luồng liên hệ / ticket (Contact)",
        "17. Luồng vận hành Admin (overview, jobs, audit)",
        "18. Ma trận tổng hợp: Ai làm gì trong từng luồng",
        "19. Trạng thái (state machines) tóm tắt",
        "20. Ghi chú kỹ thuật và ngoại lệ",
    ]
    for item in toc_items:
        add_bullet(doc, item)

    # ========== 1 ==========
    add_heading_custom(doc, "1. Tổng quan hệ thống và mô hình phân quyền", 1)
    add_para(
        doc,
        "QBankCTU là ứng dụng web quản lý ngân hàng câu hỏi, hỗ trợ giảng viên sinh câu hỏi "
        "tự động từ tài liệu bằng LLM kết hợp RAG, rồi đưa câu hỏi qua quy trình kiểm duyệt "
        "chặt chẽ trước khi đưa vào đề thi hoặc xuất bản sang Moodle.",
    )
    add_para(
        doc,
        "Hệ thống không dùng mô hình “một quyền = một vai trò” cứng nhắc. Thay vào đó áp dụng "
        "mô hình hybrid RBAC:",
        bold=True,
    )
    add_bullet(doc, "Mỗi vai trò (Admin / Teacher / Reviewer) có bộ quyền mặc định.")
    add_bullet(
        doc,
        "Admin luôn có toàn bộ quyền hệ thống; các quyền admin.* không thể gán cho vai trò khác.",
    )
    add_bullet(
        doc,
        "Teacher và Reviewer có thể được THÊM (permission_grants) hoặc BỎ (permission_revokes) "
        "từng quyền nghiệp vụ so với mặc định — chỉ trong phạm vi quyền được phép gán (ASSIGNABLE).",
    )
    add_bullet(
        doc,
        "Công thức quyền hiệu lực: "
        "effective = (mặc định vai trò ∪ grants ∪ permissions cũ) − revokes; "
        "với non-Admin còn bị cắt bởi ASSIGNABLE_PERMISSIONS.",
    )
    add_para(
        doc,
        "Luồng nghiệp vụ end-to-end điển hình:",
        bold=True,
    )
    add_para(
        doc,
        "Đăng nhập → Upload tài liệu → OCR/Extract → Chunk + Index (RAG) → "
        "Sinh câu hỏi (LLM) → Giảng viên rà soát / gửi duyệt → AI Evaluation → "
        "Người duyệt (Reviewer) kiểm duyệt → APPROVED → Làm đề thi / Xuất Moodle.",
        italic=True,
    )

    # ========== 2 ==========
    add_heading_custom(doc, "2. Ba vai trò người dùng và quyền mặc định", 1)
    add_table(
        doc,
        ["Vai trò", "Ý nghĩa", "Trang mặc định sau đăng nhập", "Quyền mặc định chính"],
        [
            [
                "Admin",
                "Quản trị toàn hệ thống: người dùng, catalog AI, jobs, audit, Moodle.",
                "/tong-quan",
                "Toàn bộ quyền (admin.* + nghiệp vụ)",
            ],
            [
                "Teacher (Giảng viên)",
                "Tác giả nội dung: upload tài liệu, sinh/soạn câu hỏi, làm đề thi, quản lý học phần của mình.",
                "/sinh-cau-hoi",
                "documents.manage_own, questions.generate, questions.manage_own, share/use shared bank, exams.manage_own, catalog.subjects.manage_own",
            ],
            [
                "Reviewer (Người duyệt)",
                "Kiểm duyệt hàng đợi câu hỏi và xuất bản câu đã duyệt lên Moodle.",
                "/kiem-duyet",
                "reviews.manage, questions.export_moodle",
            ],
        ],
    )
    add_para(
        doc,
        "Lưu ý: Hệ thống không có vai trò Student. Đăng ký công khai luôn tạo tài khoản Teacher "
        "với trạng thái chờ Admin kích hoạt (PENDING).",
        italic=True,
    )
    add_para(doc, "Các cờ tài khoản độc lập với vai trò:", bold=True)
    add_bullet(doc, "is_active: nếu false → bị khóa, không đăng nhập/API được.")
    add_bullet(
        doc,
        "approval_status: PENDING | APPROVED. Tài khoản PENDING bị chặn mọi API xác thực "
        "với thông báo chờ quản trị viên kích hoạt.",
    )
    add_bullet(
        doc,
        "review_subject_ids: danh sách môn mà Reviewer phụ trách (rỗng = mọi môn). "
        "Dùng cho auto-assign và thông báo.",
    )

    # ========== 3 ==========
    add_heading_custom(doc, "3. Danh mục quyền (permission catalog)", 1)
    add_heading_custom(doc, "3.1. Quyền nghiệp vụ Giảng viên (có thể gán)", 2)
    add_table(
        doc,
        ["Mã quyền", "Nhãn hiển thị", "Cho phép làm gì"],
        [
            ["documents.manage_own", "Quản lý tài liệu của mình", "Upload, OCR, sửa page, chunk/index, quản lý tài liệu mình sở hữu"],
            ["questions.generate", "Sinh câu hỏi bằng AI", "Gọi pipeline generation (RAG + LLM)"],
            ["questions.manage_own", "Soạn và quản lý câu hỏi của mình", "CRUD câu hỏi của mình, gửi duyệt, import/export"],
            ["questions.share_bank", "Chia sẻ câu hỏi và tài liệu", "Chia sẻ câu hỏi/tài liệu cho user hoặc theo môn"],
            ["questions.use_shared_bank", "Dùng ngân hàng được chia sẻ", "Đọc/dùng câu hỏi và tài liệu được share"],
            ["exams.manage_own", "Làm đề thi", "Tạo/sửa/xuất đề thi của mình"],
            ["catalog.subjects.manage_own", "Tạo và sửa học phần của mình", "CRUD học phần/chương/CLO do mình tạo"],
        ],
    )
    add_heading_custom(doc, "3.2. Quyền kiểm duyệt (có thể gán)", 2)
    add_table(
        doc,
        ["Mã quyền", "Nhãn hiển thị", "Cho phép làm gì"],
        [
            ["reviews.manage", "Kiểm duyệt câu hỏi", "Hàng đợi duyệt, claim, quyết định APPROVE/REJECT/NEEDS_REVISION, AI evaluate"],
            ["questions.export_moodle", "Đưa câu đã duyệt lên Moodle", "Publish câu APPROVED sang Moodle target"],
        ],
    )
    add_heading_custom(doc, "3.3. Quyền chỉ dành cho Admin", 2)
    add_table(
        doc,
        ["Mã quyền", "Nhãn hiển thị", "Cho phép làm gì"],
        [
            ["admin.overview", "Xem tổng quan", "Dashboard vận hành"],
            ["admin.users", "Quản lý người dùng", "CRUD user, invite, import, phân quyền"],
            ["admin.catalog", "Quản lý học phần và cấu hình AI", "Models, prompts, evaluation policies, subjects toàn cục"],
            ["admin.audit", "Xem nhật ký", "Audit logs"],
            ["admin.jobs", "Quản lý tác vụ", "Retry/cancel jobs OCR/generation/evaluation"],
            ["admin.moodle", "Quản lý Moodle", "Targets, publications, retry"],
            ["documents.manage_all", "Quản lý mọi tài liệu", "Truy cập/sửa mọi document"],
            ["questions.manage_all", "Quản lý mọi câu hỏi", "Truy cập/sửa mọi question"],
        ],
    )

    # ========== 4 ==========
    add_heading_custom(doc, "4. Luồng xác thực, đăng ký và kích hoạt tài khoản", 1)
    add_heading_custom(doc, "4.1. Các bước", 2)
    add_bullet(doc, "Đăng ký công khai (POST /api/v1/auth/register): luôn tạo role Teacher; approval_status=PENDING; is_active=false.")
    add_bullet(doc, "Hệ thống gửi thông báo user.pending_approval tới Admin.")
    add_bullet(doc, "Admin vào /quan-ly-nguoi-dung, kích hoạt tài khoản → is_active=true và approval_status=APPROVED.")
    add_bullet(doc, "Đăng nhập Firebase (email/Google) hoặc demo mode (admin/teacher/reviewer).")
    add_bullet(doc, "Mọi API xác thực kiểm tra: token hợp lệ → user tồn tại → không PENDING → is_active.")
    add_bullet(doc, "Đăng xuất / cập nhật hồ sơ: mọi user đã xác thực.")

    add_heading_custom(doc, "4.2. Phân quyền theo bước", 2)
    add_table(
        doc,
        ["Bước", "Ai được làm", "Ghi chú"],
        [
            ["Đăng ký tài khoản", "Công khai", "Luôn ra Teacher + PENDING"],
            ["Đăng nhập", "User đã APPROVED + active", "PENDING/khóa → 403"],
            ["Kích hoạt / phê duyệt tài khoản", "Admin (admin.users)", "Bật is_active tự APPROVED"],
            ["Đổi hồ sơ / avatar", "Chính chủ (đã đăng nhập)", ""],
            ["Xem/sửa quyền grants-revokes", "Admin", "Không gán được admin.* cho Teacher/Reviewer"],
        ],
    )

    # ========== 5 ==========
    add_heading_custom(doc, "5. Luồng quản lý tài liệu (Upload → OCR → Chunk → RAG)", 1)
    add_para(
        doc,
        "Tài liệu là nguồn tri thức cho RAG. Chỉ khi document READY (artifact OK + OCR/EXTRACT xong "
        "+ chunk set + embedding INDEXED) mới dùng để sinh câu hỏi ổn định.",
    )
    add_heading_custom(doc, "5.1. Vòng đời document", 2)
    add_para(
        doc,
        "UPLOADED → PROCESSING → READY\n"
        "                    ↘ FAILED ⇄ PROCESSING (retry)\n"
        "READY / FAILED → ARCHIVED\n"
        "READY → PROCESSING (OCR/chunk/index lại)",
        italic=True,
    )
    add_heading_custom(doc, "5.2. Các bước chi tiết", 2)
    add_bullet(doc, "Upload PDF/DOC/DOCX/MD/TXT (≤50MB), có thể gắn subject_id / chapter_id.")
    add_bullet(doc, "Lưu artifact (local/S3), tạo job OCR/EXTRACT QUEUED.")
    add_bullet(doc, "Worker: PDF text layer → extract; PDF scan → EasyOCR/Docling; DOCX → text/pages; quality gate.")
    add_bullet(doc, "Teacher/Admin có thể xem/sửa cleaned_text từng trang trước khi chunk.")
    add_bullet(doc, "Chunk (recursive/semantic) → lưu chunk_sets / document_chunks.")
    add_bullet(doc, "Embedding → ChromaDB; trạng thái embedding: PENDING|INDEXING|INDEXED|FAILED|STALE.")
    add_bullet(doc, "Sau chunk thành công: kích hoạt dictionary auto-learning (background).")
    add_bullet(doc, "Retry / cancel job; reindex khi đổi embedding model; Admin có thể transfer owner.")

    add_heading_custom(doc, "5.3. Phân quyền", 2)
    add_table(
        doc,
        ["Hành động", "Teacher", "Reviewer", "Admin"],
        [
            ["Upload / OCR / sửa page / chunk / index tài liệu của mình", "✓ (documents.manage_own)", "✗ (router chặn role)", "✓"],
            ["Quản lý mọi tài liệu", "✗", "✗", "✓ (documents.manage_all)"],
            ["Chia sẻ tài liệu", "✓ (questions.share_bank + owner)", "✗ mặc định", "✓"],
            ["Đọc tài liệu được chia sẻ", "✓ (questions.use_shared_bank)", "Có thể nếu được grant", "✓"],
            ["Chuyển owner tài liệu", "✗", "✗", "✓"],
            ["Retry/cancel job tài liệu (Admin Jobs)", "Một phần trên UI tài liệu của mình", "✗", "✓ (admin.jobs)"],
        ],
    )
    add_para(
        doc,
        "Lưu ý kỹ thuật: Router OCR/RAG còn gắn Depends(require_teacher_or_admin) — "
        "Reviewer bị chặn ở tầng role dù được grant documents.manage_own.",
        italic=True,
    )

    # ========== 6 ==========
    add_heading_custom(doc, "6. Luồng sinh câu hỏi bằng AI", 1)
    add_para(
        doc,
        "Trang /sinh-cau-hoi. AI chỉ tạo bản nháp trong kết quả job; chưa vào ngân hàng chính thức "
        "cho đến khi Giảng viên lưu/gửi duyệt.",
    )
    add_heading_custom(doc, "6.1. Các bước", 2)
    add_bullet(doc, "Teacher chọn nguồn: upload mới hoặc tài liệu READY đã có.")
    add_bullet(doc, "Cấu hình ma trận sinh (question_plan): loại câu, Bloom, độ khó, số câu (thường 1–7); instruction / target_heading.")
    add_bullet(doc, "Có thể lưu/áp dụng generation preset theo tài khoản.")
    add_bullet(doc, "POST /generation/questions → job queued (có thể kèm idempotency key).")
    add_bullet(doc, "Worker: RAG retrieve chunk → ghép prompt (system + Bloom + loại câu + ngữ cảnh) → gọi LLM → post-process/validate/dedup.")
    add_bullet(doc, "Job completed/failed; FE poll/SSE theo job_id.")
    add_bullet(doc, "Teacher xem preview kết quả trên UI, chỉnh sửa local trước khi persist.")

    add_heading_custom(doc, "6.2. Phân quyền", 2)
    add_table(
        doc,
        ["Hành động", "Teacher", "Reviewer", "Admin"],
        [
            ["Gọi API sinh câu hỏi", "✓ (questions.generate)", "✗ (router Teacher|Admin)", "✓*"],
            ["Xem/hủy job generation của mình", "✓", "✗", "✓ / manage_all"],
            ["Lưu generation preset", "✓ (questions.generate)", "✗", "✓"],
            ["Truy cập trang /sinh-cau-hoi", "✓", "✗", "Theo workspace Admin (role gate)"],
        ],
    )
    add_para(
        doc,
        "*Admin có quyền backend; UI generation thiên về Teacher. Reviewer không sinh câu hỏi mặc định.",
        italic=True,
    )

    # ========== 7 ==========
    add_heading_custom(doc, "7. Luồng soạn thảo, gửi duyệt và quản lý ngân hàng câu hỏi", 1)
    add_para(
        doc,
        "Sau khi AI sinh bản nháp (hoặc Teacher soạn tay / import), câu hỏi được persist với "
        "review_status=DRAFT. Giảng viên rà soát nội dung rồi gửi vào hàng đợi kiểm duyệt.",
    )
    add_heading_custom(doc, "7.1. Các bước", 2)
    add_bullet(doc, "Persist: POST /questions → DRAFT (origin AI|MANUAL|IMPORT); evaluation_status=NOT_STARTED.")
    add_bullet(doc, "Sửa: PATCH /questions/{id} (optimistic lock expected_version) → version mới; reset evaluation; giữ DRAFT/NEEDS_REVISION hoặc cần duyệt lại nếu đã từng quyết định.")
    add_bullet(doc, "Gửi duyệt: POST /questions/{id}/submit-review — chỉ từ DRAFT hoặc NEEDS_REVISION → PENDING.")
    add_bullet(doc, "Bulk: gửi duyệt hàng loạt, sửa metadata, archive, import/export GIFT/XML/CSV/XLSX, nhân bản.")
    add_bullet(doc, "Chia sẻ ngân hàng: share_bank cho user hoặc SUBJECT scope.")
    add_bullet(doc, "Xóa mềm: lifecycle_status ACTIVE → ARCHIVED.")

    add_heading_custom(doc, "7.2. Phân quyền — ví dụ chi tiết (luồng tạo & chuẩn bị câu hỏi)", 2)
    add_para(
        doc,
        "Trong luồng tạo câu hỏi:",
        bold=True,
    )
    add_bullet(
        doc,
        "Giảng viên (Teacher) có quyền tạo câu hỏi (sinh bằng AI hoặc soạn tay), chỉnh sửa bản nháp của mình, "
        "và gửi câu hỏi vào hàng đợi kiểm duyệt sau khi đã tự rà soát nội dung.",
    )
    add_bullet(
        doc,
        "Hệ thống AI đánh giá (AI Evaluation) chạy trên hàng đợi kiểm duyệt: chấm faithfulness, relevancy, "
        "Bloom/CLO — kết quả PASSED/FAILED chỉ mang tính hỗ trợ, không tự duyệt câu hỏi.",
    )
    add_bullet(
        doc,
        "Người duyệt (Reviewer) — hoặc Admin / Teacher được grant reviews.manage — mới có quyền quyết định "
        "duyệt (APPROVED), yêu cầu sửa (NEEDS_REVISION) hoặc từ chối (REJECTED) sau khi đã xem kết quả AI "
        "(và có thể override với lý do nếu AI FAILED).",
    )
    add_bullet(
        doc,
        "Teacher mặc định không tự duyệt câu hỏi của mình; quy tắc chống conflict-of-interest: "
        "không review câu mình tạo/sửa (trừ Admin có self_review_reason).",
    )

    add_table(
        doc,
        ["Hành động", "Teacher", "Reviewer", "Admin"],
        [
            ["Tạo / sửa / xóa câu hỏi của mình", "✓ (questions.manage_own)", "✗ mặc định", "✓ (questions.manage_all)"],
            ["Gửi duyệt (submit-review)", "✓ owner", "✗", "✓"],
            ["Đọc câu hỏi shared", "✓ (use_shared_bank)", "Có thể nếu grant", "✓"],
            ["Đọc hàng đợi để duyệt", "✗ mặc định", "✓ (reviews.manage)", "✓"],
            ["Import/export định dạng file", "✓ câu của mình", "✗", "✓"],
            ["Chia sẻ câu hỏi", "✓ (share_bank)", "✗ mặc định", "✓"],
        ],
    )

    # ========== 8 ==========
    add_heading_custom(doc, "8. Luồng đánh giá AI (AI Evaluation)", 1)
    add_para(
        doc,
        "AI Evaluation là bước hỗ trợ kiểm duyệt, độc lập với review_status. "
        "Nguyên tắc cứng: AI không bao giờ tự đặt review_status = APPROVED.",
        bold=True,
    )
    add_heading_custom(doc, "8.1. Các bước", 2)
    add_bullet(doc, "Reviewer/Admin trigger: POST .../evaluations/auto → evaluation_jobs QUEUED.")
    add_bullet(doc, "Worker claim → PROCESSING trên đúng current_version_id (CAS).")
    add_bullet(doc, "Model đánh giá theo policy: faithfulness, contextual relevancy, answer relevancy, bloom/CLO → điểm, màu RED/YELLOW/GREEN, evidence, feedback.")
    add_bullet(doc, "Kết quả: PASSED hoặc FAILED. Lỗi model → ERROR (retry hoặc duyệt tay).")
    add_bullet(doc, "Sửa nội dung sau evaluate → STALE / NOT_STARTED → phải evaluate lại.")

    add_heading_custom(doc, "8.2. Phân quyền", 2)
    add_table(
        doc,
        ["Hành động", "Teacher", "Reviewer", "Admin"],
        [
            ["Enqueue / xem AI evaluation trên hàng đợi", "✗ mặc định", "✓ (reviews.manage)", "✓"],
            ["Retry evaluation job", "✗", "Một phần qua review UI", "✓ (admin.jobs)"],
            ["Tự động đổi review_status dựa trên AI", "Không ai — hệ thống không cho phép", "", ""],
        ],
    )

    # ========== 9 ==========
    add_heading_custom(doc, "9. Luồng kiểm duyệt người (Human Review)", 1)
    add_para(
        doc,
        "Đây là bước quyết định chất lượng cuối cùng trước khi câu hỏi vào đề thi hoặc Moodle. "
        "Trang /kiem-duyet.",
    )
    add_heading_custom(doc, "9.1. Phân công & claim", 2)
    add_bullet(doc, "Claim: Reviewer nhận câu PENDING → lock IN_REVIEW.")
    add_bullet(doc, "Release / Renew: trả queue hoặc gia hạn lock.")
    add_bullet(doc, "Assign thủ công / Auto-assign: Admin (và API auto theo môn, max load, prefer/strict).")
    add_bullet(doc, "Primary reviewer không làm secondary của cùng câu; không tự duyệt câu mình tạo (trừ Admin + lý do).")

    add_heading_custom(doc, "9.2. Quyết định duyệt", 2)
    add_bullet(doc, "Quyết định: APPROVED | NEEDS_REVISION | REJECTED — phải giữ lock; expected_version khớp.")
    add_bullet(doc, "Form: checklist, criterion assessments, overall note; NEEDS_REVISION bắt buộc revision_issues; REJECT cần lý do.")
    add_bullet(doc, "Approve khi AI FAILED: bắt buộc override + lý do.")
    add_bullet(doc, "APPROVED → gắn approved_version_id.")
    add_bullet(doc, "NEEDS_REVISION → Teacher sửa → submit lại → PENDING (+ thông báo RESUBMITTED).")
    add_bullet(doc, "Secondary review (theo policy): override AI / điểm thấp / môn trong danh sách → giữ PENDING + AWAITING_SECONDARY đến khi reviewer 2 approve.")
    add_bullet(doc, "Bulk approve giới hạn (~10 câu đạt AI/GREEN), ghi audit từng item.")

    add_heading_custom(doc, "9.3. Phân quyền", 2)
    add_table(
        doc,
        ["Hành động", "Teacher", "Reviewer", "Admin"],
        [
            ["Xem dashboard kiểm duyệt", "✗ mặc định (trừ khi được grant reviews.manage)", "✓", "✓"],
            ["Claim / release / renew", "✗ / nếu được grant", "✓", "✓ (bypass lock)"],
            ["Auto-assign / review policy", "✗", "✗", "✓ (role Admin)"],
            ["Quyết định APPROVE / NEEDS_REVISION / REJECT", "✗ mặc định", "✓ (có lock)", "✓"],
            ["Duyệt câu mình tạo", "✗ (bị chặn)", "✗ nếu là author", "✓ kèm self_review_reason"],
            ["Secondary review", "✗", "✓ (reviewer khác)", "✓"],
            ["Comment / mention trên câu", "✓ nếu đọc được câu", "✓", "✓"],
        ],
    )

    # ========== 10 ==========
    add_heading_custom(doc, "10. Luồng xuất bản Moodle", 1)
    add_heading_custom(doc, "10.1. Các bước", 2)
    add_bullet(doc, "Chỉ publish approved_version_id của câu review_status=APPROVED.")
    add_bullet(doc, "Tạo moodle_publications (idempotency key) → mock PUBLISHED hoặc REST thật → cập nhật publication_status.")
    add_bullet(doc, "NOT_PUBLISHED → QUEUED → PUBLISHED | FAILED; nội dung đổi sau publish → STALE.")
    add_bullet(doc, "Admin quản lý targets, allowed_roles, retry publications.")

    add_heading_custom(doc, "10.2. Phân quyền", 2)
    add_table(
        doc,
        ["Hành động", "Teacher", "Reviewer", "Admin"],
        [
            ["Publish câu đã duyệt", "✗ mặc định", "✓ (reviews.manage + export_moodle + role ∈ target.allowed_roles)", "✓"],
            ["Cấu hình Moodle targets", "✗", "✗", "✓ (admin.moodle)"],
            ["Retry publication thất bại", "✗", "✗ / hạn chế", "✓"],
        ],
    )

    # ========== 11 ==========
    add_heading_custom(doc, "11. Luồng làm đề thi", 1)
    add_para(doc, "Trang /lam-de-thi. Chỉ dùng câu hỏi đã APPROVED với approved_version_id khớp version hiện tại, cùng môn.")
    add_heading_custom(doc, "11.1. Vòng đời đề thi", 2)
    add_para(
        doc,
        "DRAFT ⇄ READY → FINALIZED → ARCHIVED\n"
        "Sửa nội dung/matrix khi READY → về DRAFT.\n"
        "FINALIZED không quay DRAFT/READY; chỉ FINALIZED mới ARCHIVED.",
        italic=True,
    )
    add_heading_custom(doc, "11.2. Các bước", 2)
    add_bullet(doc, "Tạo đề: tên, môn, số câu, header (trường, khoa, thời gian…).")
    add_bullet(doc, "Thiết kế ma trận: chương × Bloom × độ khó × count; xem availability.")
    add_bullet(doc, "Chọn câu: auto-pick theo matrix hoặc chọn tay từ pool APPROVED.")
    add_bullet(doc, "Snapshot nội dung vào đề → READY (đủ câu) → FINALIZED.")
    add_bullet(doc, "Tạo mã đề (variant ≤4): xáo câu/đáp án, answer key.")
    add_bullet(doc, "Preview / export PDF hoặc DOCX (đề / đáp án / đề+đáp án); nhân bản → bản nháp mới.")
    add_bullet(doc, "Nếu câu trong đề bị reopen kiểm duyệt: thông báo EXAM_QUESTION_NEEDS_REVIEW cho owner đề.")

    add_heading_custom(doc, "11.3. Phân quyền", 2)
    add_table(
        doc,
        ["Hành động", "Teacher", "Reviewer", "Admin"],
        [
            ["CRUD đề thi của mình", "✓ (exams.manage_own)", "✗", "✓ (owner hoặc Admin)"],
            ["Xem/sửa đề của người khác", "✗", "✗", "✓"],
            ["Export PDF/DOCX", "✓ đề của mình", "✗", "✓"],
        ],
    )

    # ========== 12 ==========
    add_heading_custom(doc, "12. Luồng Catalog (học phần, model, prompt, policy)", 1)
    add_heading_custom(doc, "12.1. Học phần / chương / CLO", 2)
    add_bullet(doc, "CRUD: tạo, sửa, sắp xếp, deactivate (không hard-delete); chặn trùng mã.")
    add_bullet(doc, "Teacher quản lý học phần của mình (catalog.subjects.manage_own); Admin quản lý toàn cục (admin.catalog).")
    add_bullet(doc, "Đọc danh sách môn: nhiều permission reader (generate, manage_own, reviews.manage, exams…)." )

    add_heading_custom(doc, "12.2. AI models / prompts / evaluation & review policies", 2)
    add_bullet(doc, "Model catalog: activate/deactivate, version, health-check — chỉ Admin (admin.catalog).")
    add_bullet(doc, "Prompt templates: version, activate, test-build.")
    add_bullet(doc, "Evaluation policies: ngưỡng điểm/màu.")
    add_bullet(doc, "Review policy 2 vòng: secondary_on_override, secondary_below_score, secondary_subject_ids — Admin.")

    add_heading_custom(doc, "12.3. Phân quyền", 2)
    add_table(
        doc,
        ["Hành động", "Teacher", "Reviewer", "Admin"],
        [
            ["CRUD học phần của mình", "✓", "Đọc (nếu có perm reader)", "✓ full"],
            ["Cấu hình model / prompt / evaluation policy", "✗", "✗", "✓"],
            ["Cấu hình review policy (secondary)", "✗", "Đọc policy nếu reviews.manage", "✓ ghi"],
        ],
    )

    # ========== 13 ==========
    add_heading_custom(doc, "13. Luồng Dictionary (từ điển từ khóa)", 1)
    add_para(
        doc,
        "Mục đích: đo mật độ tri thức khi chunk/RAG. Chạy chủ yếu tự động sau bước chunk.",
    )
    add_bullet(doc, "Bootstrap: CORE keywords mặc định theo course_id.")
    add_bullet(doc, "Auto-learning: lấy mẫu trang → LLM trích thuật ngữ → PENDING keywords.")
    add_bullet(doc, "Active = CORE ∪ LEARNED.")
    add_para(
        doc,
        "Phân quyền: hệ thống tự chạy nền; chưa có UI duyệt PENDING→LEARNED đầy đủ trên frontend. "
        "Không gắn trực tiếp vào Teacher/Reviewer workspace.",
        italic=True,
    )

    # ========== 14 ==========
    add_heading_custom(doc, "14. Luồng quản lý người dùng (Admin)", 1)
    add_bullet(doc, "List/filter theo role, search, approval_status.")
    add_bullet(doc, "Tạo user mọi role; invite; import CSV; reset password link.")
    add_bullet(doc, "Sửa role, permission_grants/revokes, review_subject_ids, khóa/mở.")
    add_bullet(doc, "Bảo vệ Admin cuối cùng: không deactivate/demote hết Admin.")
    add_para(doc, "Quyền bắt buộc: admin.users. Teacher/Reviewer không vào /quan-ly-nguoi-dung.", bold=True)

    # ========== 15 ==========
    add_heading_custom(doc, "15. Luồng thông báo (Notifications)", 1)
    add_para(doc, "Mọi user đăng nhập xem thông báo của chính mình; mark read / mark all.")
    add_table(
        doc,
        ["Loại thông báo", "Khi nào", "Người nhận"],
        [
            ["QUESTION_SUBMITTED_FOR_REVIEW", "Teacher gửi duyệt lần đầu", "Reviewer phù hợp môn"],
            ["QUESTION_RESUBMITTED", "Gửi lại từ NEEDS_REVISION", "Reviewer"],
            ["QUESTION_APPROVED / REJECTED / NEEDS_REVISION", "Quyết định review", "Owner câu hỏi"],
            ["QUESTION_REVIEW_ASSIGNED", "Được gán", "Reviewer"],
            ["QUESTION_SECONDARY_REVIEW_*", "Cần / có duyệt lần 2", "Secondary / liên quan"],
            ["EXAM_QUESTION_NEEDS_REVIEW", "Câu trong đề bị reopen", "Owner đề"],
            ["user.pending_approval", "Đăng ký mới", "Admin"],
            ["Contact reply/status", "Admin xử lý ticket", "Requester"],
            ["Mention trong comment", "@user", "User được mention"],
        ],
    )

    # ========== 16 ==========
    add_heading_custom(doc, "16. Luồng liên hệ / ticket (Contact)", 1)
    add_para(
        doc,
        "Status: NEW → IN_PROGRESS → RESOLVED → CLOSED; có thể WITHDRAWN khi NEW/IN_PROGRESS.",
        italic=True,
    )
    add_table(
        doc,
        ["Hành động", "Teacher / Reviewer", "Admin"],
        [
            ["Tạo ticket (REVIEW_REQUEST, BUG, SUPPORT, FEEDBACK, CONTENT_ISSUE)", "✓", "✓"],
            ["Sửa nội dung khi NEW chưa có message", "✓ owner", "✓"],
            ["Rút ticket (WITHDRAWN)", "✓ owner (NEW/IN_PROGRESS)", "✓"],
            ["Chat message", "✓ owner", "✓"],
            ["Đổi status / respond chính thức", "✗", "✓"],
            ["Xem mọi ticket", "✗ (chỉ của mình)", "✓"],
        ],
    )

    # ========== 17 ==========
    add_heading_custom(doc, "17. Luồng vận hành Admin (overview, jobs, audit)", 1)
    add_table(
        doc,
        ["Module / Trang", "Quyền", "Việc chính"],
        [
            ["/tong-quan", "admin.overview", "Job breakdown, chất lượng, publication, latency/cost"],
            ["/quan-ly-job", "admin.jobs", "List/filter/retry/cancel OCR, generation, evaluation; export"],
            ["/nhat-ky-he-thong", "admin.audit", "Đọc audit theo actor/entity/action/date"],
            ["/quan-ly-moodle", "admin.moodle", "Targets, connection check, publications"],
            ["/duyet-ai", "Admin (FE admin.ai_review)", "Góc nhìn Admin về AI review"],
            ["/lich-cong-viec", "Reviewer / Admin", "Task manual + event từ queue document/question/review"],
        ],
    )

    # ========== 18 ==========
    add_heading_custom(doc, "18. Ma trận tổng hợp: Ai làm gì trong từng luồng", 1)
    add_table(
        doc,
        ["Luồng / Việc", "Teacher", "Reviewer", "Admin"],
        [
            ["Đăng ký / chờ kích hoạt", "Tự đăng ký", "—", "Kích hoạt"],
            ["Upload & xử lý tài liệu", "✓ own", "✗", "✓ all"],
            ["Sinh câu hỏi AI", "✓", "✗", "✓*"],
            ["Soạn / gửi duyệt câu hỏi", "✓ own", "✗", "✓"],
            ["AI Evaluation", "✗", "✓", "✓"],
            ["Kiểm duyệt (approve/reject/sửa)", "✗ mặc định", "✓", "✓"],
            ["Xuất Moodle", "✗ mặc định", "✓", "✓"],
            ["Làm đề thi / export PDF", "✓ own", "✗", "✓"],
            ["Catalog học phần", "✓ own", "đọc", "✓ full + AI config"],
            ["Quản lý user / jobs / audit", "✗", "✗", "✓"],
            ["Contact ticket", "✓ tạo/theo dõi", "✓ tạo/theo dõi", "✓ xử lý"],
        ],
    )

    # ========== 19 ==========
    add_heading_custom(doc, "19. Trạng thái (state machines) tóm tắt", 1)
    add_heading_custom(doc, "19.1. Bốn trục trạng thái của câu hỏi", 2)
    add_table(
        doc,
        ["Trục", "Giá trị chính", "Ai đổi"],
        [
            ["lifecycle_status", "ACTIVE → ARCHIVED", "Owner / Admin"],
            ["evaluation_status", "NOT_STARTED → QUEUED → PROCESSING → PASSED|FAILED|ERROR|STALE", "Worker AI / khi sửa nội dung"],
            ["review_status", "DRAFT → PENDING → APPROVED|REJECTED|NEEDS_REVISION", "Teacher submit; Reviewer/Admin quyết định"],
            ["publication_status", "NOT_PUBLISHED → QUEUED → PUBLISHED|FAILED|STALE", "Publish Moodle / khi nội dung đổi"],
        ],
    )
    add_heading_custom(doc, "19.2. Review assignment", 2)
    add_para(doc, "UNASSIGNED → ASSIGNED (admin/auto) → IN_REVIEW (claim + lock) → UNASSIGNED (release/timeout).", italic=True)
    add_heading_custom(doc, "19.3. User / Exam / Contact / Document", 2)
    add_bullet(doc, "User approval: PENDING → APPROVED (kèm is_active).")
    add_bullet(doc, "Exam: DRAFT ⇄ READY → FINALIZED → ARCHIVED.")
    add_bullet(doc, "Contact: NEW → IN_PROGRESS → RESOLVED → CLOSED | WITHDRAWN.")
    add_bullet(doc, "Document: UPLOADED → PROCESSING → READY | FAILED → ARCHIVED.")

    # ========== 20 ==========
    add_heading_custom(doc, "20. Ghi chú kỹ thuật và ngoại lệ", 1)
    add_bullet(
        doc,
        "Cross-role grants: Teacher có thể được thêm reviews.manage (xuất hiện trong danh sách reviewer); "
        "Reviewer có thể được thêm quyền Teacher — nhưng router OCR/generation vẫn yêu cầu role Teacher|Admin.",
    )
    add_bullet(
        doc,
        "Frontend Admin: có thể vào trang theo role gate dù permission key FE (admin.ai_review) không nằm trong "
        "ALL_PERMISSIONS backend — Admin vẫn vào được vì kiểm tra role.",
    )
    add_bullet(
        doc,
        "Object ACL: Question/Document đọc theo owner | shared+use_shared | reviews.manage | manage_all; "
        "ghi theo owner | manage_all. Exam: owner hoặc Admin.",
    )
    add_bullet(
        doc,
        "Nguồn tham chiếu code: backend/core/dependencies.py; backend/modules/{auth,users,questions,documents,exams,generation,admin,catalog,contact,notifications}; "
        "frontend/src/auth/permissions.js; docs/DATABASE_DESIGN_V2.md.",
    )
    add_para(
        doc,
        "— Hết tài liệu —",
        italic=True,
    )
    last = doc.paragraphs[-1]
    last.alignment = WD_ALIGN_PARAGRAPH.CENTER

    return doc


def main():
    out = Path("/workspace/docs/LUONG_NGHIEP_VU_VA_PHAN_QUYEN.docx")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(out)
    print(f"Wrote {out} ({out.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
